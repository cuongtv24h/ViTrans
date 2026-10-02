"""Báo cáo (SPEC §6.8–§6.10, §13): xem, xoá mềm, xuất tệp, báo lỗi đoạn, chấm sao.

Bản Markdown trong `reports.markdown` là nguồn sự thật để xuất tệp; mọi tệp xuất đều kèm
thông báo "nội dung do AI tổng hợp" ở đầu và cuối (Luật AI 134/2025/QH15, SPEC §14.4).
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path

import psycopg
from fastapi import APIRouter, Depends, Query, Request, Response

from visynth_api.db import Database
from visynth_api.errors import Problem, sql_problem
from visynth_api.models import BlockFlagIn, FeedbackIn
from visynth_api.security import current_user

router = APIRouter()

AI_NOTICE = "Nội dung do AI tổng hợp, có thể chứa sai sót. Hãy đối chiếu nguồn cho các quyết định quan trọng."
REPORT_COLUMNS = (
    "id, job_id, document_id, title, level, stats, scope_note_md, quality_grade, version, created_at, deleted_at"
)


def _db(request: Request) -> Database:
    return request.app.state.db


def _report_or_404(db: Database, report_id: str, user: dict) -> dict:
    row = db.one(
        f"SELECT {REPORT_COLUMNS} FROM reports WHERE id = %s AND user_id = %s AND deleted_at IS NULL",
        (report_id, user["id"]),
    )
    if row is None:
        raise Problem(404, "not_found", "không có báo cáo này")
    return row


def _public(report: dict) -> dict:
    row = dict(report)
    row["ai_notice"] = AI_NOTICE
    stats = row.get("stats") or {}
    row["stats"] = {
        "coverage_core": stats.get("coverage_core"),
        "faithfulness_rate": stats.get("faithfulness_rate"),
        "flagged_blocks": stats.get("flagged_blocks", 0),
        "source_words": stats.get("source_words"),
        "report_words": stats.get("report_words"),
    }
    return row


@router.get("/reports", tags=["reports"])
def list_reports(
    request: Request,
    limit: int = 20,
    cursor: str | None = None,
    user: dict = Depends(current_user),
) -> dict:
    """Danh sách báo cáo của tôi (mới nhất trước)."""
    db = _db(request)
    limit = max(1, min(limit, 100))
    rows = db.all(
        f"""SELECT {REPORT_COLUMNS} FROM reports
             WHERE user_id = %s AND deleted_at IS NULL
               AND (%s::timestamptz IS NULL OR created_at < %s::timestamptz)
             ORDER BY created_at DESC LIMIT %s""",
        (user["id"], cursor, cursor, limit + 1),
    )
    items = [_public(row) for row in rows[:limit]]
    return {"items": items, "next_cursor": items[-1]["created_at"].isoformat() if len(rows) > limit else None}


@router.get("/reports/{report_id}", tags=["reports"])
def get_report(report_id: str, request: Request, user: dict = Depends(current_user)) -> dict:
    """Báo cáo có cấu trúc (mục, khối, trích dẫn) để hiển thị."""
    db = _db(request)
    report = _report_or_404(db, report_id, user)
    sections = db.all(
        "SELECT section_id, idx, kind, title_vi FROM report_sections WHERE report_id = %s ORDER BY idx",
        (report_id,),
    )
    blocks = db.all(
        """SELECT section_id, block_id, idx, type, markdown_vi, cites, verdict, issues, status
             FROM report_blocks WHERE report_id = %s ORDER BY section_id, idx""",
        (report_id,),
    )
    by_section: dict[str, list[dict]] = {}
    for block in blocks:
        by_section.setdefault(block["section_id"], []).append(block)
    return {
        **_public(report),
        "markdown": db.scalar("SELECT markdown FROM reports WHERE id = %s", (report_id,)),
        "sections": [dict(section, blocks=by_section.get(section["section_id"], [])) for section in sections],
        "citations": db.all(
            "SELECT block_id, cite FROM report_blocks b, unnest(b.cites) AS cite WHERE report_id = %s ORDER BY block_id, cite",
            (report_id,),
        ),
    }


@router.delete("/reports/{report_id}", status_code=204, tags=["reports"])
def delete_report(report_id: str, request: Request, user: dict = Depends(current_user)) -> Response:
    """Xoá mềm báo cáo (giữ lại để đối chiếu kiểm toán)."""
    db = _db(request)
    _report_or_404(db, report_id, user)
    db.execute("UPDATE reports SET deleted_at = now() WHERE id = %s", (report_id,))
    return Response(status_code=204)


@router.post("/reports/{report_id}/blocks/{block_id}/flag", status_code=204, tags=["reports"])
def flag_block(
    report_id: str, block_id: str, body: BlockFlagIn, request: Request, user: dict = Depends(current_user)
) -> Response:
    """Báo lỗi một đoạn (sai/thiếu/khó đọc); dữ liệu đưa vào golden set sau khi được duyệt."""
    db = _db(request)
    _report_or_404(db, report_id, user)
    exists = db.scalar(
        "SELECT block_id FROM report_blocks WHERE report_id = %s AND block_id = %s", (report_id, block_id)
    )
    if exists is None:
        raise Problem(404, "not_found", "không có đoạn này trong báo cáo")
    db.execute(
        "INSERT INTO feedback (report_id, user_id, kind, block_id, tags, comment) VALUES (%s, %s, 'block_flag', %s, %s, %s)",
        (report_id, user["id"], block_id, [body.reason], body.comment),
    )
    return Response(status_code=204)


@router.post("/reports/{report_id}/feedback", status_code=204, tags=["reports"])
def report_feedback(report_id: str, body: FeedbackIn, request: Request, user: dict = Depends(current_user)) -> Response:
    """Chấm sao và góp ý cho báo cáo."""
    db = _db(request)
    _report_or_404(db, report_id, user)
    db.execute(
        "INSERT INTO feedback (report_id, user_id, kind, rating, tags, comment) VALUES (%s, %s, 'report', %s, %s, %s)",
        (report_id, user["id"], body.rating, body.tags, body.comment),
    )
    return Response(status_code=204)


# ------------------------------------------------------------------ xuất tệp


@router.get("/reports/{report_id}/export", tags=["reports"])
def export_report(
    report_id: str,
    request: Request,
    format: str = Query(..., pattern=r"^(md|docx|html|pdf)$"),
    bilingual: bool = False,
    user: dict = Depends(current_user),
) -> Response:
    """Xuất file (cache ở bảng `exports`): `md`, `html`, `docx`; `pdf` chưa hỗ trợ ở M1."""
    db: Database = _db(request)
    report = _report_or_404(db, report_id, user)
    if bilingual:
        raise Problem(422, "bilingual_unsupported", "xuất song ngữ chỉ áp dụng cho level=full_translation (M2)")
    if format == "pdf":
        raise Problem(501, "pdf_not_available", "xuất PDF sẽ bật cùng bộ chuyển đổi ở M2; dùng md/docx/html")
    markdown = db.scalar("SELECT markdown FROM reports WHERE id = %s", (report_id,)) or ""
    if not markdown:
        raise Problem(409, "report_not_ready", "báo cáo chưa có nội dung")
    document = _with_notice(markdown, report["title"])
    if format == "md":
        payload, media, ext = document.encode("utf-8"), "text/markdown; charset=utf-8", "md"
    elif format == "html":
        payload, media, ext = _to_html(document).encode("utf-8"), "text/html; charset=utf-8", "html"
    else:
        payload, media, ext = (
            _to_docx(document),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )

    storage_dir = Path(request.app.state.settings.upload_dir) / "exports"
    storage_dir.mkdir(parents=True, exist_ok=True)
    path = storage_dir / f"{report_id}.{ext}"
    path.write_bytes(payload)
    try:
        db.execute(
            """INSERT INTO exports (report_id, format, storage_key, byte_size, expires_at)
               VALUES (%s, %s, %s, %s, now() + interval '30 days')
               ON CONFLICT (report_id, format)
               DO UPDATE SET storage_key = EXCLUDED.storage_key, byte_size = EXCLUDED.byte_size,
                             created_at = now(), expires_at = now() + interval '30 days'""",
            (report_id, format, str(path), len(payload)),
        )
    except psycopg.Error as exc:
        raise sql_problem(exc) from exc
    filename = f"{report_id}.{ext}"
    return Response(
        content=payload,
        media_type=media,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-AI-Notice": "true",  # nội dung thông báo nằm trong tệp xuất (header HTTP chỉ nhận latin-1)
            "X-Export-Created-At": datetime.now(UTC).isoformat(),
        },
    )


def _with_notice(markdown: str, title: str) -> str:
    banner = f"> ⚠️ {AI_NOTICE}\n\n"
    return f"# {title}\n\n{banner}{markdown}\n\n---\n\n{AI_NOTICE}\n"


def _to_html(markdown: str) -> str:
    """HTML tối giản — M1 chỉ cần đọc được; bản đẹp (CSS, mục lục) thuộc M2."""
    escaped = markdown.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    body = "\n".join(f"<p>{line}</p>" if line.strip() else "" for line in escaped.splitlines())
    return f'<!doctype html>\n<html lang="vi">\n<head><meta charset="utf-8"><title>Báo cáo</title></head>\n<body>\n{body}\n</body>\n</html>\n'


def _to_docx(markdown: str) -> bytes:
    """DOCX từ Markdown mức tối giản: tiêu đề `#`, đoạn văn, danh sách `-`."""
    from docx import Document as DocxDocument

    document = DocxDocument()
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = min(4, len(stripped) - len(stripped.lstrip("#")))
            document.add_heading(stripped.lstrip("#").strip(), level=level)
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(stripped)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
