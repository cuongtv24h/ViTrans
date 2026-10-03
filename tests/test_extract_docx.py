"""Bóc tách DOCX: style tiêu đề, danh sách, bảng, tracked-changes, chống zip-bomb — SPEC §6.1."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from visynth.extract import ExtractionError, check_zip_safety, extract, extract_docx

docx = pytest.importorskip("docx")
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


def _make_docx(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Chương 1", level=1)
    doc.add_paragraph("Đoạn thân bài có dấu tiếng Việt.")
    doc.add_heading("Mục 1.1", level=2)
    doc.add_paragraph("Ý trong danh sách", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tên"
    table.cell(0, 1).text = "Giá"
    table.cell(1, 0).text = "A | B"  # dấu | phải được thoát
    table.cell(1, 1).text = "1"
    doc.save(str(path))


def test_docx_structure(tmp_path: Path):
    p = tmp_path / "bai.docx"
    _make_docx(p)
    ext = extract_docx(p)
    assert [x.kind for x in ext.paragraphs] == ["heading", "body", "heading", "list_item", "table"]
    assert ext.paragraphs[0].content == "Chương 1"
    assert [s.title for s in ext.sections] == ["Chương 1", "Mục 1.1"]
    assert ext.sections[1].parent_section_id == "S01"
    # bảng → Markdown có dòng phân cách và thoát dấu |
    table_md = ext.paragraphs[-1].content
    assert table_md.splitlines()[0] == "| Tên | Giá |"
    assert table_md.splitlines()[1] == "|---|---|"
    assert "A \\| B" in table_md
    assert ext.title == "bai"


def test_docx_tracked_changes_keep_final_version(tmp_path: Path):
    p = tmp_path / "tracked.docx"
    doc = docx.Document()
    para = doc.add_paragraph()

    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "1")
    ins.set(qn("w:author"), "tester")
    r_ins = OxmlElement("w:r")
    t_ins = OxmlElement("w:t")
    t_ins.text = "bản cuối"
    r_ins.append(t_ins)
    ins.append(r_ins)

    dele = OxmlElement("w:del")
    dele.set(qn("w:id"), "2")
    r_del = OxmlElement("w:r")
    t_del = OxmlElement("w:delText")
    t_del.text = "bản cũ"
    r_del.append(t_del)
    dele.append(r_del)

    para._p.append(ins)
    para._p.append(dele)
    doc.save(str(p))

    ext = extract_docx(p)
    assert ext.paragraphs[0].content == "bản cuối"


def test_images_ignored_warning(tmp_path: Path):
    import base64
    import io

    png_1x1 = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    p = tmp_path / "anh.docx"
    doc = docx.Document()
    doc.add_paragraph("Có ảnh ở đây.")
    doc.add_picture(io.BytesIO(png_1x1))
    doc.save(str(p))
    ext = extract_docx(p)
    assert "images_ignored" in ext.warnings


def test_zip_bomb_ratio_is_rejected(tmp_path: Path):
    bomb = tmp_path / "bomb.zip"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("big.txt", "0" * 5_000_000)
    with zipfile.ZipFile(bomb) as zf, pytest.raises(ExtractionError) as excinfo:
        check_zip_safety(zf)
    assert excinfo.value.code in ("file_too_large", "extraction_failed")


def test_non_docx_zip_is_rejected(tmp_path: Path):
    other = tmp_path / "khac.docx"
    with zipfile.ZipFile(other, "w") as zf:
        zf.writestr("hello.txt", "không phải DOCX")
    with pytest.raises(ExtractionError) as excinfo:
        extract_docx(other)
    assert excinfo.value.code == "unsupported_file_type"


def test_broken_pdf_and_oversize_are_rejected(tmp_path: Path):
    with pytest.raises(ExtractionError) as e1:
        extract(b"%PDF-1.7 fake pdf", title="x")
    assert e1.value.code == "unreadable_pdf"

    with pytest.raises(ExtractionError) as e2:
        extract(b"x" * 5000, title="x", max_bytes=1000)
    assert e2.value.code == "file_too_large"


def test_extract_dispatch_by_magic_bytes(tmp_path: Path):
    p = tmp_path / "bai.docx"
    _make_docx(p)
    ext = extract(p)
    assert ext.title == "bai"
    assert ext.source_type == "upload"

    txt = tmp_path / "ghi_chu.md"
    txt.write_text("# Ghi chú\n\nNội dung.", encoding="utf-8")
    ext2 = extract(txt)
    assert ext2.paragraphs[0].kind == "heading"
